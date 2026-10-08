"""분석 결과 + 콘셉트 후보 -> 신제품_콘셉트후보_YYYYMMDD.xlsx"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# 점수 가중치(합 1.0). 상품팀 기준에 맞게 바꿔도 된다.
WEIGHTS = {"need": 0.30, "sales": 0.25, "trend": 0.15, "differentiation": 0.15, "feasibility": 0.15}
SCORE_LABELS = {"need": "고객니즈", "sales": "매출연관", "trend": "트렌드", "differentiation": "차별성", "feasibility": "OEM가능"}

GREEN = "2E7D32"
HEAD = PatternFill("solid", fgColor=GREEN)
HEAD_FONT = Font(bold=True, color="FFFFFF")
REVIEW_FILL = PatternFill("solid", fgColor="FFF8E1")
WARN_FILL = PatternFill("solid", fgColor="FFEBEE")
THIN = Border(*(Side(style="thin", color="DDDDDD"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
DISCLAIMER = "※ 점수와 순위는 AI가 데이터로 낸 참고 의견입니다. 최종 선택은 상품팀이 합니다. AI는 사람과 다른 기준(예: 가격 비중을 더 크게 봄)으로 고르는 경향이 있으니 근거 리뷰를 꼭 직접 확인해 주세요."


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font, cell.border = HEAD, HEAD_FONT, THIN
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def put(ws, row, values, fill=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.alignment, c.border = WRAP, THIN
        if fill:
            c.fill = fill


def quote(r):
    star = f"★{r['rating']:g} " if r.get("rating") is not None else ""
    return f"[{r['id']}] {star}\"{r['text']}\" — {r['file']} {r['row']}행"


def build_excel(a, cj, out_dir):
    """엑셀을 저장하고 (파일 경로, 근거 확인 필요 경고 목록, 점수순으로 정렬된 콘셉트 목록)을 돌려준다."""
    concepts = cj["concepts"] if isinstance(cj, dict) else cj
    rv = {r["id"]: r for r in a["reviews"]}
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    warnings = []
    for c in concepts:
        sc = c.get("scores", {})
        vals = {k: max(1, min(5, float(sc.get(k, 1) or 1))) for k in WEIGHTS}
        c["_scores"] = vals
        c["_total"] = round(sum(vals[k] * w for k, w in WEIGHTS.items()) / 5 * 100, 1)
        ids = c.get("evidence_review_ids") or []
        c["_valid"] = [i for i in ids if i in rv]
        c["_invalid"] = [i for i in ids if i not in rv]
        if not c["_valid"] or c["_invalid"]:
            warnings.append(f"- {c.get('name')}: 유효 근거 {len(c['_valid'])}개, 없는 ID {c['_invalid'] or '없음'}")
    concepts.sort(key=lambda c: (-bool(c["_valid"]), -c["_total"]))

    wb = Workbook()

    # 1. 읽어주세요
    ws = wb.active
    ws.title = "읽어주세요"
    o = a["overview"]
    lines = [
        ("스윗밸런스 신제품 콘셉트 후보", None),
        (DISCLAIMER, None),
        ("", None),
        ("생성일", cj.get("generated_at", date.today().isoformat()) if isinstance(cj, dict) else date.today().isoformat()),
        ("분석 리뷰 수", o["review_count"]),
        ("리뷰 기간", " ~ ".join(o["review_period"]) if o.get("review_period") else "날짜 정보 없음"),
        ("제품 수", o["product_count"]),
        ("매출 데이터", "있음" if o["has_sales"] else "없음"),
        ("트렌드 데이터", (f"{len(a['trends'])}개 (" + ("웹 검색, 출처는 '트렌드' 시트" if any(t.get("from_web") for t in a["trends"]) else "파일") + ")") if a["trends"] else "없음"),
        ("카테고리 필터", a.get("category_filter") or "없음"),
        ("콘셉트 수", len(concepts)),
        ("메모", cj.get("notes", "") if isinstance(cj, dict) else ""),
        ("", None),
        ("총점 계산", "각 항목 1~5점 × 가중치 합산, 100점 만점"),
    ] + [(f"  {SCORE_LABELS[k]}", f"{int(w * 100)}%") for k, w in WEIGHTS.items()] + [
        ("", None),
        ("상품팀 검토 방법", "'콘셉트 후보' 시트 오른쪽 검토 칸에 채택/보류/제외, 검토자, 메모를 적어 주세요."),
        ("근거 확인 필요", "빨간색 행은 근거 리뷰 ID가 데이터에 없어서 검증이 안 된 후보입니다."),
        ("분석 파일", ", ".join(f["file"] for f in a["files"])),
    ]
    for i, (k, v) in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=k).alignment = WRAP
        if v is not None:
            ws.cell(row=i, column=2, value=v).alignment = WRAP
    ws["A1"].font = Font(bold=True, size=16, color=GREEN)
    ws["A2"].font = Font(bold=True, color="C62828")
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 48
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90

    # 2. 콘셉트 후보
    ws = wb.create_sheet("콘셉트 후보")
    ws.cell(row=1, column=1, value=DISCLAIMER).font = Font(bold=True, color="C62828")
    cols = ["순위", "콘셉트명", "카테고리", "한 줄 설명", "타깃", "핵심 재료", "가격대", "발상 방식", "근거·논리",
            "근거 리뷰 원문", "근거 수", "관련 제품", "관련 제품 달성률", "트렌드 키워드", "트렌드 출처", "리스크"] + \
           [SCORE_LABELS[k] for k in WEIGHTS] + ["총점", "검토: 결정", "검토자", "검토 메모"]
    widths = [6, 24, 10, 34, 22, 24, 12, 12, 34, 70, 7, 18, 14, 16, 26, 26] + [8] * len(WEIGHTS) + [8, 11, 10, 26]
    header(ws, 2, cols, widths)
    prod = {p["product"]: p for p in a["products"]}
    for n, c in enumerate(concepts, 1):
        rel = c.get("related_products") or []
        ach = ", ".join(f"{p} {prod[p]['achievement_pct']}%" for p in rel if p in prod and prod[p].get("achievement_pct") is not None)
        ev = "\n\n".join(quote(rv[i]) for i in c["_valid"][:4])
        if len(c["_valid"]) > 4:
            ev += f"\n\n외 {len(c['_valid']) - 4}건은 '근거 리뷰' 시트 참고"
        if c["_invalid"]:
            ev += f"\n\n[근거 확인 필요] 데이터에 없는 ID: {', '.join(c['_invalid'])}"
        row = [n, c.get("name"), c.get("category"), c.get("one_liner"), c.get("target"), c.get("key_ingredients"),
               c.get("price_range"), c.get("approach"), c.get("rationale"), ev, len(c["_valid"]), ", ".join(rel), ach,
               ", ".join(c.get("trend_keywords") or []), c.get("trend_source", ""), c.get("risks")] + \
              [c["_scores"][k] for k in WEIGHTS] + [c["_total"], "", "", ""]
        r = n + 2
        put(ws, r, row, WARN_FILL if (not c["_valid"] or c["_invalid"]) else None)
        ws.cell(row=r, column=10).fill = REVIEW_FILL if c["_valid"] and not c["_invalid"] else WARN_FILL
        ws.cell(row=r, column=2).font = Font(bold=True)
    last = len(concepts) + 2
    dec_col = get_column_letter(len(cols) - 2)
    dv = DataValidation(type="list", formula1='"채택,보류,제외"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{dec_col}3:{dec_col}{max(last, 3)}")
    tot_col = get_column_letter(len(cols) - 3)
    ws.conditional_formatting.add(f"{tot_col}3:{tot_col}{max(last, 3)}", CellIsRule(operator=">=", formula=["80"], fill=PatternFill("solid", fgColor="C8E6C9")))
    ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{max(last, 2)}"

    # 3. 근거 리뷰
    ws = wb.create_sheet("근거 리뷰")
    header(ws, 1, ["콘셉트 순위", "콘셉트명", "리뷰 ID", "제품", "평점", "작성일", "리뷰 원문", "파일", "행"], [9, 24, 9, 20, 6, 11, 80, 28, 6])
    r = 2
    for n, c in enumerate(concepts, 1):
        for i in c["_valid"]:
            x = rv[i]
            put(ws, r, [n, c.get("name"), i, x["product"], x["rating"], x.get("date"), x["text"], x["file"], x["row"]])
            r += 1

    # 4. 제품별 현황
    ws = wb.create_sheet("제품별 현황")
    header(ws, 1, ["제품", "리뷰 수", "평균 평점", "저평점 수", "매출", "목표", "달성률(%)", "자주 언급 속성", "부정 많은 속성", "재료", "가격", "kcal"],
           [24, 8, 9, 9, 14, 14, 10, 30, 26, 30, 10, 8])
    for i, p in enumerate(a["products"], 2):
        m = p.get("master") or {}
        put(ws, i, [p["product"], p["review_count"], p["avg_rating"], p["low_rating_count"], p.get("revenue"), p.get("target"),
                    p.get("achievement_pct"), ", ".join(f"{k}({v})" for k, v in p["top_attributes"]),
                    ", ".join(f"{k}({v})" for k, v in p["top_negative_attributes"]), m.get("ingredients"), m.get("price"), m.get("kcal")])
        for col in (5, 6):
            ws.cell(row=i, column=col).number_format = "#,##0"

    # 5. 속성 분석
    ws = wb.create_sheet("속성 분석")
    header(ws, 1, ["속성", "언급", "긍정", "부정", "부정 비율", "대표 부정 리뷰", "대표 긍정 리뷰"], [14, 8, 8, 8, 9, 70, 70])
    for i, (k, v) in enumerate(sorted(a["attributes"].items(), key=lambda x: -x[1]["mentions"]), 2):
        put(ws, i, [k, v["mentions"], v["positive"], v["negative"],
                    round(v["negative"] / v["mentions"], 2) if v["mentions"] else None,
                    "\n\n".join(quote(rv[x]) for x in v["negative_examples"][:3]),
                    "\n\n".join(quote(rv[x]) for x in v["positive_examples"][:3])])
        ws.cell(row=i, column=5).number_format = "0%"

    # 6. 트렌드
    ws = wb.create_sheet("트렌드")
    header(ws, 1, ["키워드", "수치", "증감", "리뷰 언급", "출처", "사용한 콘셉트"], [18, 10, 10, 10, 50, 50])
    used = {}
    for n, c in enumerate(concepts, 1):
        for k in c.get("trend_keywords") or []:
            used.setdefault(k, {"concepts": [], "sources": set()})
            used[k]["concepts"].append(f"{n}. {c.get('name')}")
            if c.get("trend_source"):
                used[k]["sources"].add(c["trend_source"])
    r = 2
    seen = set()
    for t in a["trends"]:
        u = used.get(t["keyword"], {"concepts": [], "sources": set()})
        put(ws, r, [t["keyword"], t["volume"], t["growth"], t["review_mentions"], ", ".join(t["sources"]), "\n".join(u["concepts"])])
        seen.add(t["keyword"])
        r += 1
    for k, u in used.items():
        if k not in seen:
            put(ws, r, [k, None, None, sum(1 for x in a["reviews"] if k in x["text"]),
                        ", ".join(sorted(u["sources"])) or "출처 없음(확인 필요)", "\n".join(u["concepts"])])
            r += 1

    fname = out_dir / f"신제품_콘셉트후보_{date.today().strftime('%Y%m%d')}.xlsx"
    wb.save(fname)
    return fname, warnings, concepts
